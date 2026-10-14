from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).parent
OUT=ROOT/'论文与作品集'
data=json.loads((OUT/'工作文件'/'manuscript.json').read_text(encoding='utf-8'))
for lang in ['EN','ZH']:
    z=lang=='ZH';lk='zh' if z else 'en';pk='chinese' if z else 'english'
    d=Document();sec=d.sections[0]
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.88)
    norm=d.styles['Normal'];norm.font.name='Microsoft YaHei' if z else 'Georgia'
    norm.font.size=Pt(10.5 if z else 11)
    norm.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    norm.paragraph_format.line_spacing=1.35
    norm.paragraph_format.space_after=Pt(9)
    for sty,size in [('Title',24),('Subtitle',13),('Heading 1',17),('Heading 2',12)]:
        st=d.styles[sty];st.font.name='Microsoft YaHei' if z else 'Arial'
        st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        st.font.size=Pt(20 if z and sty=='Title' else size);st.font.color.rgb=RGBColor.from_string('000000' if sty=='Title' else '345A60')
        st.paragraph_format.space_after=Pt(13)
    p=sec.header.paragraphs[0]
    p.text='桂阳县苦志育才公益学校校园研究' if z else 'KUZHI YUCAI  |  CAMPUS RESEARCH'
    p.runs[0].font.size=Pt(8);p.runs[0].font.color.rgb=RGBColor.from_string('65736D')
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    d.add_paragraph('桂阳县苦志育才公益学校校园调研与设计研究' if z else 'Campus Research and Design for Kuzhi Yucai Public Welfare School','Title')
    d.add_paragraph(data['title_'+lk],style='Subtitle')
    d.add_paragraph('2026年9月调研资料' if z else 'Field material collected in September 2026')
    for i,s in enumerate(data['sections']):
        d.add_paragraph(s[lk].replace(' / ','  '),style='Heading 1')
        for p in s[pk]:d.add_paragraph(p)
        if s['fig'] in ['keywords','audit','time']:
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(OUT/'图表与图纸'/f'{s["fig"]}_{lang}.png'),width=Inches(6.1))
            cap={
              'keywords':('Selected keyword occurrences in the Chinese questionnaire answers','中文问卷答案中选定关键词的出现次数'),
              'audit':('Raw environmental audit means on the original 1 to 5 form','环境审计原表1至5分的各维度原始均值'),
              'time':('Minutes in the school timetable between 06:30 and 21:00','6:30至21:00作息安排的分类分钟数')
            }[s['fig']]
            cp=d.add_paragraph(cap[1 if z else 0]);cp.runs[0].font.size=Pt(9)
    refheading=d.add_paragraph('参考资料' if z else 'References',style='Heading 1')
    if z:refheading.paragraph_format.page_break_before=True
    for r in data['references']:d.add_paragraph(f'[{r["id"]}] '+r[lk])
    d.core_properties.title=data['title_'+lk]
    d.core_properties.author='Student field research'
    d.save(OUT/f'Paper_{lang}_Editable.docx')
    check=OUT/'工作文件'/f'Paper_{lang}_Word_check.pdf'
    if check.exists():check.unlink()
    print('Saved',lang)
