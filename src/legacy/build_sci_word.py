from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

R=Path(__file__).parent
ROOT=R/'论文与作品集'
B=ROOT/'工作文件/SCI_Word'
OUT=ROOT/'英文Word论文'
B.mkdir(parents=True,exist_ok=True);OUT.mkdir(exist_ok=True)
M=json.loads((ROOT/'工作文件/manuscript_publication.json').read_text('utf8'))
L=json.loads((ROOT/'工作文件/ppt_build/publication_layout_EN.json').read_text('utf8'))
clouds=json.loads((ROOT/'工作文件/ppt_build/wordcloud_audit.json').read_text('utf8'))
FONT='C:/Windows/Fonts/arialbd.ttf'
# Three faithful word clouds, composed into one vertically aligned figure.
im=Image.new('RGB',(1800,2710),'white');draw=ImageDraw.Draw(im)
for k,g in enumerate(clouds):
    y0=k*890
    draw.text((40,y0+22),f'({chr(97+k)}) {g["title"]}',fill='#252D2C',font=ImageFont.truetype(FONT,38))
    scale=6.3
    offset=(1800-g['width']*scale)/2
    top=y0+91+(750-g['height']*scale)/2
    for a in g['words']:
        f=ImageFont.truetype(FONT,round(a['size']*scale))
        draw.text((offset+a['x']*scale,top+a['y']*scale),a['english'],font=f,fill=a['color'])
im.save(B/'Fig1_WordClouds.png',dpi=(300,300))
# Quantitative results share one clean, consistent figure.
plt.rcParams.update({'font.family':'Arial','font.size':10,'pdf.fonttype':42})
fig=plt.figure(figsize=(7.2,5.8))
ax=fig.add_axes([.25,.44,.69,.46])
labels=['Basic safety','Comfort','Activity support','Care support','Spatial connections']
vals=[3.75,3.44,2.71,3.83,3.00]
bars=ax.barh(labels[::-1],vals[::-1],height=.60,color=['#81968A','#345A60','#345A60','#81968A','#81968A'])
ax.bar_label(bars,fmt='%.2f',padding=5)
ax.set_xlim(0,5);ax.set_xticks(range(6));ax.set_xlabel('Mean score')
ax.xaxis.grid(True,color='#E5EAE7',linewidth=.5);ax.set_axisbelow(True)
for sp in ax.spines.values():sp.set_visible(False)
ax.tick_params(length=0)
fig.text(.035,.953,'(a) Environmental assessment',weight='bold',fontsize=12)
fig.text(.035,.325,'(b) Allocation of daily time',weight='bold',fontsize=12)
ax2=fig.add_axes([.06,.20,.88,.07]);offset=0
names=['Study','Meals','Activity','Midday rest','Eye exercises','Intervals']
minutes=[390,160,90,80,30,120]
colors=['#345A60','#81968A','#687F72','#AAB8AA','#BDC8BE','#D9DFD9']
for name,v,color in zip(names,minutes,colors):
    ax2.barh([0],[v],left=offset,color=color,label=name,height=.7)
    ax2.text(offset+v/2,0,str(v),ha='center',va='center',fontsize=9,color='white' if name=='Study' else '#252D2C')
    offset+=v
ax2.set_xlim(0,870);ax2.set_axis_off()
ax2.legend(loc='upper center',bbox_to_anchor=(.5,-.24),ncol=3,frameon=False,fontsize=9)
fig.savefig(B/'Fig2_Environment_Time.png',dpi=320,facecolor='white');plt.close(fig)
# Architectural proposal in a two-panel figure.
plan=Image.open(L['slides'][14]['objects'][8]['path']).convert('RGBA')
aerial=Image.open(L['slides'][21]['objects'][8]['path']).convert('RGBA')
im=Image.new('RGB',(2100,1350),'white');dr=ImageDraw.Draw(im)
for source,box,title in [(plan,(20,95,1000,1220),'(a) Campus masterplan'),(aerial,(1070,200,1000,1000),'(b) Connected wings and courtyard')]:
    x,y,w,h=box;source.thumbnail((w,h),Image.Resampling.LANCZOS)
    im.paste(source,(x+(w-source.width)//2,y+(h-source.height)//2),source)
    dr.text((x,25),title,fill='#252D2C',font=ImageFont.truetype(FONT,35))
im.save(B/'Fig3_Campus.png',dpi=(300,300))

d=Document();sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.3)
sec.left_margin=sec.right_margin=Cm(2.5)
sec.header_distance=sec.footer_distance=Cm(1.1)
for st in d.styles:
    if st.type==1:
        st.font.name='Times New Roman'
        st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
    for bd in list(st.element.iter(qn('w:pBdr'))):bd.getparent().remove(bd)
norm=d.styles['Normal'];norm.font.size=Pt(12)
norm.paragraph_format.line_spacing=1.5
norm.paragraph_format.space_after=Pt(6)
norm.paragraph_format.widow_control=True
for name,size in [('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
    st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(12 if name!='Title' else 0)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.1
    st.paragraph_format.keep_with_next=True
cap=d.styles['Caption'];cap.font.name='Times New Roman';cap.font.size=Pt(10)
cap.font.italic=False;cap.font.bold=False;cap.paragraph_format.line_spacing=1.15
cap.paragraph_format.space_after=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
title='Everyday Care and Campus Design at Kuzhi Yucai Public Welfare School'
d.add_paragraph(title,'Title')

bodytexts=[]
def p(text):
    para=d.add_paragraph(text)
    para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    bodytexts.append(text);return para
def heading(text,level=1):return d.add_paragraph(text,f'Heading {level}')
def sectionparas(idx,indices=None):
    arr=M['sections'][idx]['english']
    for k in indices if indices is not None else range(len(arr)):p(arr[k])
def caption(text):
    para=d.add_paragraph(text,'Caption')
    return para
def picture(path,width):
    para=d.add_paragraph()
    para.paragraph_format.line_spacing=1
    para.paragraph_format.space_after=Pt(3)
    para.paragraph_format.keep_with_next=True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path),width=Inches(width))

heading('Abstract')
# Merge the existing abstract into the customary continuous journal paragraph.
p(' '.join(M['sections'][0]['english']))
kw=d.add_paragraph();kw.paragraph_format.line_spacing=1.15
kw.add_run('Keywords: ').bold=True
kw.add_run('School design; everyday care; activity settings; environmental audit; rural campus')
heading('1 Introduction');sectionparas(1)
heading('2 Materials and methods')
heading('2.1 Study setting',2);sectionparas(7,[0,1,2])
heading('2.2 Fieldwork and questionnaire analysis',2);sectionparas(2,[0])
p('The field survey and school interviews were conducted in September 2026 [3,4]. Questionnaire feedback was organized into ten topics covering school contact, children’s needs, everyday use, spatial difficulties, preferred places, underused areas, improvement priorities, and resources. The complete thematic comparison appears in Appendix A.')
heading('2.3 Environmental assessment',2);sectionparas(2,[1])
heading('2.4 Keyword analysis and word clouds',2);sectionparas(2,[2])
p('Literal occurrences of selected Chinese terms were counted in the written answers, excluding question prompts. Compound terms were retained, including teaching building and school clinic. Three word clouds present recurring terms, campus spaces, and care and resources. English labels translate the Chinese terms after counting. Within each cloud, font size is proportional to the square root of frequency. Each term appears once in a cloud, and terms can recur across the three thematic views. Color distinguishes words; font size represents frequency.')
heading('2.5 Observation and spatial synthesis',2);sectionparas(2,[3])
p('The analysis combined descriptive counts, environmental scores, observed activities, and interview themes. These findings were linked to room relationships, circulation, shared spaces, and implementation priorities in the campus proposal [5].')
heading('3 Results')
heading('3.1 Questionnaire findings and recurring themes',2);sectionparas(3)
p('The word clouds bring together the vocabulary of school life (Fig. 1). Classroom and teaching building dominate the spatial terms, while funding is prominent in the care and resources view. Words such as daily care, emotional support, teaching staff, and donations connect the physical setting to the people and resources that sustain its use.')
heading('3.2 Environmental conditions',2);sectionparas(4)
p('The domain means and the daily time distribution are shown in Fig. 2. Appendix B lists all environmental assessment items and their recorded scores.')
heading('3.3 Daily rhythms and space use',2);sectionparas(5)
heading('3.4 Care and school operations',2);sectionparas(6)
heading('4 Discussion')
heading('4.1 Translating everyday needs into design priorities',2);sectionparas(8)
heading('4.2 Architectural organization and shared spaces',2);sectionparas(7,[3]);sectionparas(9)
p('The masterplan and courtyard perspective illustrate the proposed relationship between the teaching wings, shared facilities, and outdoor grounds (Fig. 3). The spatial organization brings the study’s priorities into a connected campus arrangement [5].')
heading('4.3 Implementation and follow up',2);sectionparas(10)
heading('5 Conclusion');sectionparas(11)
heading('References')
for n,r in enumerate(M['references'],1):
    para=d.add_paragraph(f'[{n}] {r["en"]}')
    para.paragraph_format.line_spacing=1.15
    para.paragraph_format.space_after=Pt(8)
    para.paragraph_format.left_indent=Cm(.65)
    para.paragraph_format.first_line_indent=Cm(-.65)
    for run in para.runs:run.font.size=Pt(10.5)

# Place full-size figure plates after the references, keeping each plate together.
d.add_page_break()
picture(B/'Fig1_WordClouds.png',5.45)
caption('Fig. 1. Word clouds of questionnaire responses. (a) Recurring terms. (b) Campus spaces. (c) Care and resources. Frequencies were counted in the Chinese answers and translated into English labels. Word size is proportional to the square root of frequency within each panel.')
d.add_page_break()
picture(B/'Fig2_Environment_Time.png',5.65)
caption('Fig. 2. Environmental assessment and the school timetable. (a) Mean scores across five environmental domains on a scale of 1–5. (b) Minutes allocated to study, meals, activity, midday rest, eye exercises, and intervals between 06:30 and 21:00.')
picture(B/'Fig3_Campus.png',5.8)
caption('Fig. 3. Campus design proposal. (a) Masterplan showing teaching and shared facilities, sports grounds, and open space. (b) Connected teaching wings frame a courtyard for daily movement and shared activity.')

def table(headers,rows,widths,size=10):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    pr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');pr.append(rep)
    for row in rows:
        cs=t.add_row().cells
        for i,val in enumerate(row):cs[i].text=str(val)
    for rno,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i])
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge=='bottom' else 'nil')
                e.set(qn('w:sz'),'4');e.set(qn('w:color'),'C6CDCA');borders.append(e)
            tcpr.append(borders)
            if rno==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFF3F0');tcpr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.line_spacing=1.12
                para.paragraph_format.space_before=Pt(2);para.paragraph_format.space_after=Pt(2)
                for run in para.runs:run.font.size=Pt(size);run.bold=rno==0
    return t

d.add_page_break();heading('Appendix A Questionnaire feedback')
caption('Table A1. Questionnaire responses organized by topic.')
obs=L['slides'][11]['objects'];rows=[]
for j,o in enumerate(obs):
    if j>=8 and o['kind']=='text' and re.fullmatch(r'\d{2}',o['text']):
        rows.append([obs[j+1]['text'],obs[j+2]['text'],obs[j+3]['text']])
table(['Topic','Feedback A','Feedback B'],rows,[3.1,6.45,6.45],10.5)
d.add_page_break();heading('Appendix B Environmental assessment')
caption('Table B1. Environmental assessment items and recorded scores on a scale of 1–5.')
obs=L['slides'][12]['objects'];rows=[]
for j,o in enumerate(obs):
    if j>=8 and o['kind']=='text' and re.fullmatch(r'\d{2}',o['text']):
        rows.append([o['text'],obs[j+1]['text'],obs[j+2]['text']])
table(['Item','Assessment question','Score'],rows,[1.3,13.2,1.5],10)

d.core_properties.title=title
d.core_properties.subject='Campus environment research and architectural design'
d.core_properties.author=''
d.core_properties.keywords='school design, everyday care, environmental audit, rural campus'
path=OUT/'Kuzhi_Yucai_Campus_Research_EN.docx';d.save(path)
alltext='\n'.join(q.text for q in d.paragraphs)
missing=[q for secm in M['sections'] for q in secm['english'] if q not in alltext]
assert not missing,missing
audit={'original_paragraphs_retained':46,'main_text_words':len(re.findall(r"\b[\w’-]+\b",' '.join(bodytexts))),
       'questionnaire_topics':10,'audit_items':35,'figure_plates':3,'language':'English','file':str(path)}
(B/'content_audit.json').write_text(json.dumps(audit,indent=2),'utf8')
print(json.dumps(audit,indent=2))
