import re,json,csv,sys
from pathlib import Path
from PIL import Image,ImageOps,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).parent
OUT=ROOT/"研究报告交付";FIG=OUT/"图表";DATA=OUT/"分析数据"

def photo(name):
 p=ROOT/"材料整理/照片"/name
 if p.suffix.lower()=='.heic':p=ROOT/"材料整理/照片预览"/(p.stem+".jpg")
 return ImageOps.exif_transpose(Image.open(p)).convert("RGB")
def board(filename,items):
 canvas=Image.new("RGB",(1800,1420),"white")
 font=ImageFont.truetype("C:/Windows/Fonts/msyh.ttc",28)
 d=ImageDraw.Draw(canvas)
 for i,(name,label) in enumerate(items):
  im=photo(name);im.thumbnail((870,635))
  x=(i%2)*900;y=(i//2)*710
  canvas.paste(im,(x+(880-im.width)//2,y+(635-im.height)//2))
  d.text((x+14,y+650),label,font=font,fill=(35,45,50))
 canvas.save(FIG/filename)
board("06_校园已有空间.png",[
 ("IMG_3732.HEIC","教学楼与有顶连廊"),
 ("IMG_3720.JPG","食堂内的就餐桌椅"),
 ("IMG_3734.HEIC","操场及边缘铺装"),
 ("IMG_3697.HEIC","带沙盘游戏室标识的空间")])
board("07_空间与外部环境.png",[
 ("IMG_3705.HEIC","室内桌椅集中堆放"),
 ("IMG_3736.HEIC","室外草地与裸土地"),
 ("IMG_3757.JPG","校园周边道路"),
 ("IMG_3770.HEIC","周边工程机械场景")])

def readcsv(p):
 with p.open(encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
text=(ROOT/"材料整理/写作源文件/研究报告正文.md").read_text(encoding='utf-8')
audio=ROOT/"材料整理/写作源文件/录音分析章节.md"
if audio.exists():text=text.replace("{{AUDIO_SECTION}}",audio.read_text(encoding='utf-8'))
else:raise RuntimeError("Audio chapter not yet written")
audit=readcsv(DATA/"环境审计全部项目.csv")
atable="| 维度 | 项目 | 分数 |\n| --- | --- | --- |\n"+"\n".join(f'| {r["维度"]} | {r["项目"]} | {r["评分"]} |' for r in audit)
text=text.replace("{{AUDIT_TABLE}}",atable)
srows=readcsv(DATA/"作息时间统计.csv")
stable="| 安排 | 开始 | 结束 | 分钟数 |\n| --- | --- | --- | --- |\n"+"\n".join(f'| {r["安排"]} | {r["开始"]} | {r["结束"]} | {r["分钟数"]} |' for r in srows)
text=text.replace("{{SCHEDULE_TABLE}}",stable)
assert "{{" not in text
(OUT/"桂阳县苦志育才公益学校调研报告.md").write_text(text,encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.83);sec.bottom_margin=Inches(.78)
sec.left_margin=Inches(.86);sec.right_margin=Inches(.86)
sec.header_distance=Inches(.35);sec.footer_distance=Inches(.35)
sec.different_first_page_header_footer=True
styles=doc.styles
for border in list(styles.element.xpath(".//w:pBdr")):
 border.getparent().remove(border)
def setfont(style,western,east,size,bold=False):
 s=styles[style];s.font.name=western;s.font.size=Pt(size);s.font.bold=bold
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
 s.font.color.rgb=RGBColor(0,0,0)
 s.font.italic=False
setfont('Normal','Times New Roman','宋体',11.5)
normal=styles['Normal'].paragraph_format
normal.line_spacing=1.32;normal.space_after=Pt(6);normal.first_line_indent=Pt(23)
for name,size in [('Title',26),('Subtitle',14),('Heading 1',17),('Heading 2',13),('Heading 3',12)]:
 setfont(name,'Arial','黑体',size,name not in ['Title','Subtitle'])
 styles[name].paragraph_format.first_line_indent=Pt(0)
 styles[name].paragraph_format.space_before=Pt(14 if name=='Heading 1' else 10)
 styles[name].paragraph_format.space_after=Pt(8)
 styles[name].paragraph_format.keep_with_next=True
setfont('Caption','Arial','宋体',9.5)
styles['Caption'].paragraph_format.first_line_indent=Pt(0)
styles['Caption'].paragraph_format.space_after=Pt(9)
styles['Caption'].paragraph_format.line_spacing=1.15
toc_style=styles.add_style("ReportContentsHeading",WD_STYLE_TYPE.PARAGRAPH)
toc_style.base_style=styles['Normal']
setfont("ReportContentsHeading","Arial","黑体",17,True)
toc_style.paragraph_format.first_line_indent=Pt(0)
toc_style.paragraph_format.space_after=Pt(14)

header=sec.header.paragraphs[0]
header.text="桂阳县苦志育才公益学校儿童照护与校园空间调研"
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.first_line_indent=Pt(0)
for r in header.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string("666666")
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
for literal,field in [('第 ',None),('',"PAGE"),(' 页',None)]:
 run=footer.add_run(literal);run.font.size=Pt(9)
 if field:
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),field);run._r.addnext(fld)

title=doc.add_paragraph("桂阳县苦志育才公益学校\n儿童照护与校园空间\n调研报告",'Title')
title.paragraph_format.space_before=Pt(86)
title.paragraph_format.space_after=Pt(26)
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub=doc.add_paragraph("高中生社会调查研究报告",'Subtitle');sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
for line in ["调研资料时间  2026年9月23日至26日","整理时间  2026年10月"]:
 p=doc.add_paragraph(line);p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph("从问卷、访谈和现场记录出发\n观察儿童在校园里怎样学习、生活和活动")
p.paragraph_format.space_before=Pt(65);p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
doc.add_paragraph("目录",'ReportContentsHeading')
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u')
p._p.append(fld)
doc.add_page_break()

def shade(cell,color):
 pr=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),color);pr.append(s)
def table(lines):
 parsed=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
 parsed=[r for r in parsed if not all(re.fullmatch(r":?-+:?",c.replace(' ','')) for c in r)]
 if not parsed:return
 cols=len(parsed[0]);t=doc.add_table(rows=1,cols=cols)
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 usable=6.78
 # Make descriptive columns wider than numeric columns.
 if cols==3:ratios=[.22,.62,.16] if parsed[0][-1] in ["分数","评分"] else [.23,.38,.39]
 elif cols==4:ratios=[.16,.28,.28,.28]
 elif cols==5:ratios=[.15,.13,.16,.18,.38] if "主要语境" in parsed[0] else [.20,.20,.20,.20,.20]
 else:ratios=[1/cols]*cols
 if parsed[0]==["维度","项目数","总分","平均分","1至2分项目数"]:ratios=[.27,.15,.14,.18,.26]
 if parsed[0]==["较弱项目","原始评分","结合其他材料的理解"]:ratios=[.25,.14,.61]
 if parsed[0]==["文件","约略时间","本报告使用的主题"]:ratios=[.43,.23,.34]
 if parsed[0]==["日间安排","分钟数","占870分钟的比例","计算口径"]:ratios=[.23,.12,.21,.44]
 if parsed[0]==["安排","开始","结束","分钟数"]:ratios=[.4,.2,.2,.2]
 for c,ratio in zip(t.columns,ratios):c.width=Inches(usable*ratio)
 for ri,row in enumerate(parsed):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  for ci,value in enumerate(row):
   cells[ci].width=Inches(usable*ratios[ci])
   cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cells[ci].paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
   p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(14)
   snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
   if len(parsed)<=20 and ri<len(parsed)-1:p.paragraph_format.keep_with_next=True
   r=p.add_run(value);r.font.size=Pt(10)
   if ri==0:r.bold=True;shade(cells[ci],"EAF0F3")
   elif ri%2==0:shade(cells[ci],"F7F9FA")
   tcPr=cells[ci]._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    m=OxmlElement('w:'+side);m.set(qn('w:w'),'65' if side in ['top','bottom'] else '85');m.set(qn('w:type'),'dxa');mar.append(m)
   tcPr.append(mar)
  pr=t.rows[ri]._tr.get_or_add_trPr()
  cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
 # subtle horizontal lines only
 borders=OxmlElement('w:tblBorders')
 for side in ['top','bottom','insideH']:
  b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D6DEE3');borders.append(b)
 t._tbl.tblPr.append(borders)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)

lines=text.splitlines();i=0;body=False
while i<len(lines):
 line=lines[i].strip()
 if line=="## 摘要":body=True
 if not body or not line:i+=1;continue
 if line.startswith('|'):
  ts=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   ts.append(lines[i]);i+=1
  table(ts);continue
 if line.startswith('## '):
  heading=line[3:];p=doc.add_paragraph(heading,'Heading 1')
  # Keep the long report organized, with appendices each starting cleanly.
  if heading.startswith('附录'):p.paragraph_format.page_break_before=True
 elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.keep_with_next=True
  fp=OUT/m.group(2);im=Image.open(fp)
  width=min(6.65,4.95*im.width/im.height)
  p.add_run().add_picture(str(fp),width=Inches(width))
 elif re.match(r'^图\d+\s',line):
  p=doc.add_paragraph(line,'Caption')
 elif line.startswith('关键词'):
  p=doc.add_paragraph(line);p.paragraph_format.first_line_indent=Pt(0)
  p.runs[0].bold=True
 else:doc.add_paragraph(line)
 i+=1

# Enable automatic field refresh in Word.
settings=doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
doc.core_properties.title="桂阳县苦志育才公益学校儿童照护与校园空间调研报告"
doc.core_properties.subject="学生调研材料整理与分析"
doc.core_properties.author="调研小组"
doc.core_properties.keywords="儿童照护 校园空间 问卷 词频 环境审计"
path=OUT/"桂阳县苦志育才公益学校调研报告.docx"
doc.save(path)
print(path)
print("Chinese characters",len(re.findall(r"[\u4e00-\u9fff]",text)))
